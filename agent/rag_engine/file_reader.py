import json
import re
from pathlib import Path
from typing import Any

from config import DATA_DIRS, SUPPORTED_EXTENSIONS
from rag_engine.path_filter import iter_source_files

try:
    from docx import Document
    from pypdf import PdfReader
except ModuleNotFoundError:
    Document = None
    PdfReader = None


def read_text_file(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp950"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue

    return path.read_text(encoding="utf-8", errors="replace")


def read_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    parts: list[str] = []

    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text.strip())

    return "\n\n".join(parts)


def read_docx(path: Path) -> str:
    document = Document(path)
    parts: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def read_ipynb(path: Path) -> str:
    notebook = json.loads(path.read_text(encoding="utf-8"))
    parts: list[str] = []

    for cell in notebook.get("cells", []):
        if cell.get("cell_type") not in {"markdown", "code"}:
            continue

        source = cell.get("source", "")
        if isinstance(source, list):
            text = "".join(source)
        else:
            text = str(source)

        if text.strip():
            parts.append(text.strip())

    return "\n\n".join(parts)


def _base_metadata(path: Path) -> dict[str, Any]:
    metadata: dict[str, Any] = {
        "source_path": str(path),
        "file_name": path.name,
        "file_type": path.suffix.lower(),
        "folder_name": path.parent.name,
    }

    for root in DATA_DIRS:
        try:
            relative_path = path.resolve().relative_to(root.resolve())
        except (OSError, ValueError):
            continue

        relative_parts = relative_path.parts
        relative_text = str(relative_path)
        week_match = re.search(r"(?i)week[\s_-]?(\d+)", relative_text)
        metadata.update(
            {
                "root_source": root.name,
                "relative_path": relative_text,
                "project": root.parent.name if root.name.lower() == "learning" else root.name,
                "module": relative_parts[0] if len(relative_parts) > 1 else path.stem,
                "week": f"Week{int(week_match.group(1)):02d}" if week_match else "",
            }
        )
        break

    metadata.setdefault("root_source", path.parent.name)
    metadata.setdefault("relative_path", path.name)
    metadata.setdefault("project", path.parent.name)
    metadata.setdefault("module", path.stem)
    metadata.setdefault("week", "")
    return metadata


def _read_pdf_pages(path: Path, metadata: dict[str, Any]) -> list[dict[str, Any]]:
    reader = PdfReader(str(path))
    docs: list[dict[str, Any]] = []

    total_pages = len(reader.pages)
    extracted_pages = 0
    total_characters = 0

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            extracted_pages += 1
            total_characters += len(text.strip())
            docs.append(
                {
                    "path": path,
                    "text": text.strip(),
                    "metadata": {
                        **metadata,
                        "page_number": page_number,
                        "pdf_total_pages": total_pages,
                    },
                }
            )

    if not docs:
        docs.append(
            {
                "path": path,
                "text": "",
                "metadata": {**metadata, "pdf_total_pages": total_pages},
            }
        )

    empty_pages = total_pages - extracted_pages
    average_characters = total_characters / extracted_pages if extracted_pages else 0.0
    print(f"PDF 擷取品質：{path}")
    print(f"總頁數：{total_pages}")
    print(f"成功擷取頁數：{extracted_pages}")
    print(f"空白頁數：{empty_pages}")
    print(f"平均每個有效頁面字數：{average_characters:.1f}")
    if total_pages and extracted_pages / total_pages < 0.5:
        print("警告：可擷取文字的頁面比例偏低，這份 PDF 可能需要 OCR。")

    return docs


def read_files(folders: list[Path]) -> list[dict]:
    docs: list[dict] = []

    for folder in folders:
        if not folder.exists():
            print(f"找不到資料夾：{folder}")
            continue

        for path in iter_source_files(folder):
            if not path.is_file():
                continue

            suffix = path.suffix.lower()
            if suffix not in SUPPORTED_EXTENSIONS:
                continue

            print(f"正在處理：{path}")
            metadata = _base_metadata(path)

            try:
                if suffix == ".pdf":
                    docs.extend(_read_pdf_pages(path, metadata))
                elif suffix == ".docx":
                    docs.append({"path": path, "text": read_docx(path), "metadata": metadata})
                elif suffix == ".ipynb":
                    docs.append({"path": path, "text": read_ipynb(path), "metadata": metadata})
                else:
                    docs.append({"path": path, "text": read_text_file(path), "metadata": metadata})
            except Exception as exc:
                print(f"讀取失敗：{path}")
                print(f"原因：{exc}")
                docs.append(
                    {
                        "path": path,
                        "text": "",
                        "metadata": {
                            **metadata,
                            "read_error": str(exc),
                        },
                    }
                )

    return docs


def read_single_file(path: Path) -> list[dict]:
    metadata = _base_metadata(path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return _read_pdf_pages(path, metadata)
    if suffix == ".docx":
        return [{"path": path, "text": read_docx(path), "metadata": metadata}]
    if suffix == ".ipynb":
        return [{"path": path, "text": read_ipynb(path), "metadata": metadata}]

    return [{"path": path, "text": read_text_file(path), "metadata": metadata}]
