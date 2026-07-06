# Legacy single-file version. New entry point is main.py.

from __future__ import annotations

import ast
import hashlib
import json
import os
import re
import sys
import time
from collections import deque
from datetime import datetime
from pathlib import Path
from typing import Any

SKIP_MODEL_LOAD = os.environ.get("RAG_AGENT_SKIP_MODEL_LOAD") == "1"

try:
    import chromadb
    import requests
    from docx import Document
    from pypdf import PdfReader
    from sentence_transformers import CrossEncoder, SentenceTransformer
except ModuleNotFoundError:
    if not SKIP_MODEL_LOAD:
        raise
    chromadb = None
    requests = None
    Document = None
    PdfReader = None
    CrossEncoder = None
    SentenceTransformer = None

OLLAMA_URL = "http://localhost:11434/api/generate"
MODEL = "qwen2.5:14b"
# MODEL = "qwen2.5-coder:14b"

DATA_DIRS = [
    Path(r"D:\AI_Server\knowledge_base"),
    Path(r"D:\NKUST_VLM\learning"),
]

DB_DIR = r"D:\AI_Server\agent\chroma_db"
COLLECTION_NAME = "nkust_knowledge_bge_m3"
MANIFEST_PATH = Path(r"D:\AI_Server\agent\index_manifest.json")
EMBEDDING_MODEL = "BAAI/bge-m3"
RERANKER_MODEL = "BAAI/bge-reranker-v2-m3"
RERANKER_ENABLED = True
VECTOR_SEARCH_TOP_K = 20
RERANK_TOP_K = 5
CONVERSATION_MEMORY_SIZE = 5

SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".py",
    ".cpp",
    ".h",
    ".hpp",
    ".json",
    ".yaml",
    ".yml",
    ".ipynb",
    ".pdf",
    ".docx",
}

if SKIP_MODEL_LOAD:
    embedding_model = None
    reranker_model = None
else:
    try:
        print("載入 Embedding 模型：")
        print(EMBEDDING_MODEL)
        embedding_model = SentenceTransformer(EMBEDDING_MODEL)
        print("Embedding 模型載入完成。")
    except Exception as exc:
        print("Embedding 模型載入失敗：")
        print(exc)
        sys.exit(1)

    reranker_model = None
    if RERANKER_ENABLED:
        try:
            print("載入 Reranker 模型：")
            print(RERANKER_MODEL)
            reranker_model = CrossEncoder(RERANKER_MODEL)
            print("Reranker 模型載入完成。")
        except Exception as exc:
            print("Reranker 模型載入失敗：")
            print(exc)
            print("將 fallback 回原本的向量搜尋結果。")

if SKIP_MODEL_LOAD:
    client = None
    collection = None
else:
    client = chromadb.PersistentClient(path=DB_DIR)
    collection = client.get_or_create_collection(name=COLLECTION_NAME)


def encode_text(text: str) -> list[float]:
    if embedding_model is None:
        raise RuntimeError("Embedding 模型尚未載入。")

    return embedding_model.encode(
        text,
        normalize_embeddings=True,
    ).tolist()


def read_text_file(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "cp950"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue

    return path.read_text(encoding="utf-8", errors="replace")


def read_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    parts: list[str] = []

    for page in reader.pages:
        text = page.extract_text() or ""
        if text.strip():
            parts.append(text.strip())

    return "\n\n".join(parts)


def read_docx(path: Path) -> str:
    document = Document(path)
    parts: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts)


def read_ipynb(path: Path) -> str:
    notebook = json.loads(path.read_text(encoding="utf-8"))
    parts: list[str] = []

    for cell in notebook.get("cells", []):
        if cell.get("cell_type") not in {"markdown", "code"}:
            continue

        source = cell.get("source", "")
        if isinstance(source, list):
            text = "".join(source)
        else:
            text = str(source)

        if text.strip():
            parts.append(text.strip())

    return "\n\n".join(parts)


def _base_metadata(path: Path) -> dict[str, Any]:
    return {
        "source_path": str(path),
        "file_name": path.name,
        "file_type": path.suffix.lower(),
        "folder_name": path.parent.name,
    }


def _read_pdf_pages(path: Path, metadata: dict[str, Any]) -> list[dict[str, Any]]:
    reader = PdfReader(str(path))
    docs: list[dict[str, Any]] = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            docs.append(
                {
                    "path": path,
                    "text": text.strip(),
                    "metadata": {
                        **metadata,
                        "page_number": page_number,
                    },
                }
            )

    if not docs:
        docs.append({"path": path, "text": "", "metadata": metadata})

    return docs


def read_files(folders: list[Path]) -> list[dict]:
    docs: list[dict] = []

    for folder in folders:
        if not folder.exists():
            print(f"找不到資料夾：{folder}")
            continue

        for path in folder.rglob("*"):
            if not path.is_file():
                continue

            suffix = path.suffix.lower()
            if suffix not in SUPPORTED_EXTENSIONS:
                continue

            print(f"正在處理：{path}")
            metadata = _base_metadata(path)

            try:
                if suffix == ".pdf":
                    docs.extend(_read_pdf_pages(path, metadata))
                elif suffix == ".docx":
                    docs.append({"path": path, "text": read_docx(path), "metadata": metadata})
                elif suffix == ".ipynb":
                    docs.append({"path": path, "text": read_ipynb(path), "metadata": metadata})
                else:
                    docs.append({"path": path, "text": read_text_file(path), "metadata": metadata})
            except Exception as exc:
                print(f"讀取失敗：{path}")
                print(f"原因：{exc}")
                docs.append(
                    {
                        "path": path,
                        "text": "",
                        "metadata": {
                            **metadata,
                            "read_error": str(exc),
                        },
                    }
                )

    return docs


def calculate_file_sha256(path: Path) -> str:
    sha256 = hashlib.sha256()
    with path.open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            sha256.update(block)
    return sha256.hexdigest()


def create_empty_manifest() -> dict:
    return {
        "embedding_model": EMBEDDING_MODEL,
        "collection_name": COLLECTION_NAME,
        "data_dirs": [str(path) for path in DATA_DIRS],
        "files": {},
    }


def load_manifest() -> dict:
    if not MANIFEST_PATH.exists():
        return create_empty_manifest()

    try:
        with MANIFEST_PATH.open("r", encoding="utf-8") as file:
            manifest = json.load(file)
    except Exception as exc:
        backup_path = MANIFEST_PATH.with_name("index_manifest.bak.json")
        try:
            backup_path.write_text(MANIFEST_PATH.read_text(encoding="utf-8"), encoding="utf-8")
            print(f"manifest 損壞，已備份至：{backup_path}")
        except Exception as backup_exc:
            print(f"manifest 損壞，且備份失敗：{backup_exc}")
        print(f"將建立新的 manifest。原因：{exc}")
        return create_empty_manifest()

    if not isinstance(manifest, dict):
        print("manifest 格式不正確，將建立新的 manifest。")
        return create_empty_manifest()

    manifest.setdefault("embedding_model", EMBEDDING_MODEL)
    manifest.setdefault("collection_name", COLLECTION_NAME)
    manifest.setdefault("data_dirs", [str(path) for path in DATA_DIRS])
    manifest.setdefault("files", {})
    return manifest


def save_manifest(manifest: dict) -> None:
    MANIFEST_PATH.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def scan_source_files(folders: list[Path]) -> list[Path]:
    files: list[Path] = []

    for folder in folders:
        if not folder.exists():
            print(f"找不到資料夾：{folder}")
            continue

        for path in folder.rglob("*"):
            if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS:
                files.append(path)

    return sorted(files, key=lambda item: str(item).lower())


def delete_file_chunks(chunk_ids: list[str]) -> None:
    if not chunk_ids:
        return

    try:
        collection.delete(ids=chunk_ids)
    except Exception as exc:
        print(f"刪除舊 chunks 時發生警告：{exc}")


def read_single_file(path: Path) -> list[dict]:
    metadata = _base_metadata(path)
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        return _read_pdf_pages(path, metadata)
    if suffix == ".docx":
        return [{"path": path, "text": read_docx(path), "metadata": metadata}]
    if suffix == ".ipynb":
        return [{"path": path, "text": read_ipynb(path), "metadata": metadata}]

    return [{"path": path, "text": read_text_file(path), "metadata": metadata}]


def split_text(text: str, chunk_size: int = 800, overlap: int = 150) -> list[str]:
    text = text.strip()
    if not text:
        return []

    if overlap >= chunk_size:
        raise ValueError("overlap 必須小於 chunk_size")

    chunks: list[str] = []
    start = 0

    while start < len(text):
        end = start + chunk_size
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)

        if end >= len(text):
            break

        start = end - overlap

    return chunks


def _split_large_chunk(chunk: str, max_size: int = 1500) -> list[str]:
    if len(chunk) <= max_size:
        return [chunk.strip()] if chunk.strip() else []

    return split_text(chunk)


def split_markdown(text: str) -> list[str]:
    text = text.strip()
    if not text:
        return []

    heading_pattern = re.compile(r"(?m)^(#{1,4})\s+(.+?)\s*$")
    matches = list(heading_pattern.finditer(text))
    if not matches:
        return split_text(text)

    chunks: list[str] = []
    leading_text = text[: matches[0].start()].strip()
    if leading_text:
        chunks.extend(_split_large_chunk(leading_text))

    for index, match in enumerate(matches):
        start = match.start()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        section = text[start:end].strip()
        chunks.extend(_split_large_chunk(section))

    return chunks


def split_python(text: str) -> list[str]:
    text = text.strip()
    if not text:
        return []

    try:
        tree = ast.parse(text)
    except SyntaxError:
        return split_text(text)

    lines = text.splitlines()
    nodes = [
        node
        for node in tree.body
        if isinstance(node, (ast.ClassDef, ast.FunctionDef, ast.AsyncFunctionDef))
        and hasattr(node, "end_lineno")
    ]

    if not nodes:
        return split_text(text)

    chunks: list[str] = []
    covered_ranges: list[tuple[int, int]] = []

    for node in nodes:
        start_line = node.lineno
        if getattr(node, "decorator_list", None):
            start_line = min(decorator.lineno for decorator in node.decorator_list)

        end_line = node.end_lineno or node.lineno
        chunk = "\n".join(lines[start_line - 1 : end_line]).strip()
        if chunk:
            chunks.append(chunk)
            covered_ranges.append((start_line, end_line))

    extra_parts: list[str] = []
    cursor = 1
    for start_line, end_line in sorted(covered_ranges):
        if cursor < start_line:
            extra = "\n".join(lines[cursor - 1 : start_line - 1]).strip()
            if extra:
                extra_parts.append(extra)
        cursor = max(cursor, end_line + 1)

    if cursor <= len(lines):
        extra = "\n".join(lines[cursor - 1 :]).strip()
        if extra:
            extra_parts.append(extra)

    for extra in extra_parts:
        chunks.extend(split_text(extra))

    return chunks


def _find_matching_brace(text: str, open_brace_index: int) -> int | None:
    depth = 0
    in_string: str | None = None
    escaped = False

    for index in range(open_brace_index, len(text)):
        char = text[index]

        if in_string:
            if escaped:
                escaped = False
            elif char == "\\":
                escaped = True
            elif char == in_string:
                in_string = None
            continue

        if char in {'"', "'"}:
            in_string = char
        elif char == "{":
            depth += 1
        elif char == "}":
            depth -= 1
            if depth == 0:
                return index

    return None


def split_cpp(text: str) -> list[str]:
    text = text.strip()
    if not text:
        return []

    pattern = re.compile(
        r"(?m)^[\t ]*(?:"
        r"(?:class|struct)\s+[A-Za-z_]\w*[^{;]*\{"
        r"|(?:[A-Za-z_][\w:<>,~*&\s]+)\s+[A-Za-z_]\w*\s*\([^;{}]*\)\s*(?:const\s*)?\{"
        r")"
    )
    matches = list(pattern.finditer(text))
    if not matches:
        return split_text(text)

    chunks: list[str] = []
    covered_ranges: list[tuple[int, int]] = []

    for match in matches:
        open_brace = text.find("{", match.start(), match.end())
        if open_brace == -1:
            continue

        close_brace = _find_matching_brace(text, open_brace)
        if close_brace is None:
            continue

        end = close_brace + 1
        if end < len(text) and text[end] == ";":
            end += 1

        chunk = text[match.start() : end].strip()
        if chunk:
            chunks.append(chunk)
            covered_ranges.append((match.start(), end))

    if not chunks:
        return split_text(text)

    extra_parts: list[str] = []
    cursor = 0
    for start, end in sorted(covered_ranges):
        if cursor < start:
            extra = text[cursor:start].strip()
            if extra:
                extra_parts.append(extra)
        cursor = max(cursor, end)

    if cursor < len(text):
        extra = text[cursor:].strip()
        if extra:
            extra_parts.append(extra)

    for extra in extra_parts:
        chunks.extend(split_text(extra))

    return chunks


def split_text_by_paragraph(text: str) -> list[str]:
    text = text.strip()
    if not text:
        return []

    paragraphs = [paragraph.strip() for paragraph in re.split(r"\n\s*\n+", text) if paragraph.strip()]
    if not paragraphs:
        return split_text(text)

    chunks: list[str] = []
    current = ""

    for paragraph in paragraphs:
        if len(paragraph) > 1200:
            if current:
                chunks.append(current.strip())
                current = ""
            chunks.extend(split_text(paragraph))
            continue

        candidate = f"{current}\n\n{paragraph}".strip() if current else paragraph
        if len(candidate) <= 1200:
            current = candidate
            continue

        if current:
            chunks.append(current.strip())
        current = paragraph

    if current:
        chunks.append(current.strip())

    balanced_chunks: list[str] = []
    buffer = ""
    for chunk in chunks:
        candidate = f"{buffer}\n\n{chunk}".strip() if buffer else chunk
        if buffer and len(buffer) < 500 and len(candidate) <= 1200:
            buffer = candidate
        else:
            if buffer:
                balanced_chunks.append(buffer)
            buffer = chunk

    if buffer:
        balanced_chunks.append(buffer)

    return balanced_chunks


def _chunk_strategy(file_type: str, text: str) -> str:
    file_type = file_type.lower()

    if file_type == ".md" and re.search(r"(?m)^#{1,4}\s+", text):
        return "markdown_heading"
    if file_type == ".py":
        try:
            ast.parse(text)
            return "python_ast"
        except SyntaxError:
            return "fixed_window"
    if file_type in {".cpp", ".h", ".hpp"}:
        if re.search(r"(?m)^[\t ]*(?:class|struct)\s+[A-Za-z_]\w*|^[\t ]*(?:[A-Za-z_][\w:<>,~*&\s]+)\s+[A-Za-z_]\w*\s*\([^;{}]*\)\s*(?:const\s*)?\{", text):
            return "cpp_regex"
        return "fixed_window"
    if file_type in {".txt", ".docx", ".pdf", ".ipynb", ".json", ".yaml", ".yml"}:
        return "paragraph"
    return "fixed_window"


def chunk_document(text: str, file_type: str) -> list[str]:
    file_type = file_type.lower()

    if file_type == ".md":
        return split_markdown(text)
    if file_type == ".py":
        return split_python(text)
    if file_type in {".cpp", ".h", ".hpp"}:
        return split_cpp(text)
    if file_type in {".txt", ".docx", ".pdf", ".ipynb", ".json", ".yaml", ".yml"}:
        return split_text_by_paragraph(text)

    return split_text(text)


def clear_collection() -> None:
    try:
        existing = collection.get()
        ids = existing.get("ids", [])
        if ids:
            collection.delete(ids=ids)
    except Exception as exc:
        print(f"清除舊索引時發生警告：{exc}")


def _markdown_section_title(text_before_chunk: str) -> str | None:
    matches = re.findall(r"(?m)^\s{0,3}#{1,6}\s+(.+?)\s*$", text_before_chunk)
    return matches[-1].strip() if matches else None


def _python_function_or_class(chunk: str) -> str | None:
    try:
        tree = ast.parse(chunk)
    except SyntaxError:
        match = re.search(r"(?m)^\s*(?:async\s+def|def|class)\s+([A-Za-z_]\w*)", chunk)
        return match.group(1) if match else None

    for node in ast.walk(tree):
        if isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef, ast.ClassDef)):
            return node.name

    return None


def _cpp_function_or_class(chunk: str) -> str | None:
    class_match = re.search(r"\b(?:class|struct)\s+([A-Za-z_]\w*)", chunk)
    if class_match:
        return class_match.group(1)

    function_match = re.search(
        r"\b(?:[A-Za-z_][\w:<>,~*&\s]+)\s+([A-Za-z_]\w*)\s*\([^;{}]*\)\s*(?:const\s*)?\{",
        chunk,
    )
    if function_match:
        return function_match.group(1)

    return None


def _chunk_metadata(
    doc: dict,
    chunk: str,
    chunk_id: str,
    start_index: int,
    chunk_strategy: str,
) -> dict[str, Any]:
    metadata = dict(doc["metadata"])
    metadata["chunk_id"] = chunk_id
    metadata["chunk_strategy"] = chunk_strategy

    file_type = metadata.get("file_type", "").lower()

    if file_type == ".md":
        section_title = _markdown_section_title(doc["text"][:start_index] + chunk)
        if section_title:
            metadata["section_title"] = section_title
    elif file_type == ".py":
        function_or_class = _python_function_or_class(chunk)
        if function_or_class:
            metadata["function_or_class"] = function_or_class
    elif file_type in {".cpp", ".h", ".hpp"}:
        function_or_class = _cpp_function_or_class(chunk)
        if function_or_class:
            metadata["function_or_class"] = function_or_class

    return metadata


def index_single_file(path: Path, file_sha256: str) -> dict | None:
    try:
        docs = read_single_file(path)
    except Exception as exc:
        print(f"讀取失敗：{path}")
        print(f"原因：{exc}")
        return None

    stat = path.stat()
    chunk_ids: list[str] = []
    chunk_index = 0

    for doc in docs:
        text = doc.get("text", "")
        if not text.strip():
            continue

        metadata = doc.get("metadata", {})
        file_type = metadata.get("file_type", "").lower()
        chunk_strategy = _chunk_strategy(file_type, text)
        chunks = chunk_document(text, file_type)
        if not chunks:
            continue

        search_start = 0

        for chunk in chunks:
            chunk_id = f"{file_sha256[:16]}_{chunk_index}"

            print("正在建立向量：")
            print(path.name)
            print(f"strategy: {chunk_strategy}")

            start_index = text.find(chunk[:80], search_start)
            if start_index == -1:
                start_index = search_start
            search_start = start_index + max(len(chunk) - 150, 1)

            embedding = encode_text(chunk)
            chunk_metadata = _chunk_metadata(doc, chunk, chunk_id, start_index, chunk_strategy)

            collection.add(
                ids=[chunk_id],
                documents=[chunk],
                embeddings=[embedding],
                metadatas=[chunk_metadata],
            )

            chunk_ids.append(chunk_id)
            chunk_index += 1

    if not chunk_ids:
        print(f"檔案沒有可索引內容，已略過：{path}")
        return None

    return {
        "source_path": str(path),
        "sha256": file_sha256,
        "mtime": stat.st_mtime,
        "file_size": stat.st_size,
        "indexed_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "chunk_ids": chunk_ids,
    }


def build_index(full_rebuild: bool = False) -> None:
    start_time = time.perf_counter()

    if full_rebuild:
        clear_collection()
        manifest = create_empty_manifest()
        source_files = scan_source_files(DATA_DIRS)
        indexed_files = 0
        skipped_files = 0
        chunk_count = 0

        for path in source_files:
            try:
                file_sha256 = calculate_file_sha256(path)
                entry = index_single_file(path, file_sha256)
            except Exception as exc:
                print(f"索引檔案失敗：{path}")
                print(f"原因：{exc}")
                entry = None

            if entry is None:
                skipped_files += 1
                continue

            manifest["files"][str(path)] = entry
            indexed_files += 1
            chunk_count += len(entry.get("chunk_ids", []))

        save_manifest(manifest)
        elapsed = time.perf_counter() - start_time

        print("\n完整索引重建完成")
        print(f"索引檔案數：{indexed_files}")
        print(f"略過檔案數：{skipped_files}")
        print(f"Chunks：{chunk_count}")
        print(f"耗時：{elapsed:.1f} 秒")
        return

    manifest_exists = MANIFEST_PATH.exists()
    if not manifest_exists:
        try:
            existing_count = collection.count()
        except Exception:
            existing_count = 0

        if existing_count > 0:
            print("偵測到尚未建立 manifest，但 ChromaDB 可能已有舊索引。")
            print("建議先執行「完整重建索引」。")
            return

    manifest = load_manifest()

    if manifest.get("embedding_model") != EMBEDDING_MODEL:
        print("目前 Embedding 模型與 manifest 不一致。")
        print("建議執行完整重建索引。")
        return

    if manifest.get("collection_name") != COLLECTION_NAME:
        print("目前 ChromaDB collection 與 manifest 不一致。")
        print("建議執行完整重建索引。")
        return

    source_files = scan_source_files(DATA_DIRS)
    current_paths = {str(path) for path in source_files}
    manifest_files = manifest.setdefault("files", {})

    new_file_count = 0
    updated_file_count = 0
    deleted_file_count = 0
    skipped_unchanged_count = 0
    changed_chunk_count = 0

    for source_path in list(manifest_files.keys()):
        if source_path not in current_paths:
            entry = manifest_files[source_path]
            delete_file_chunks(entry.get("chunk_ids", []))
            del manifest_files[source_path]
            deleted_file_count += 1

    for path in source_files:
        path_key = str(path)

        try:
            file_sha256 = calculate_file_sha256(path)
        except Exception as exc:
            print(f"計算檔案 sha256 失敗：{path}")
            print(f"原因：{exc}")
            continue

        old_entry = manifest_files.get(path_key)
        if old_entry and old_entry.get("sha256") == file_sha256:
            skipped_unchanged_count += 1
            continue

        is_update = old_entry is not None
        if old_entry:
            delete_file_chunks(old_entry.get("chunk_ids", []))

        try:
            entry = index_single_file(path, file_sha256)
        except Exception as exc:
            print(f"索引檔案失敗：{path}")
            print(f"原因：{exc}")
            entry = None

        if entry is None:
            if is_update:
                manifest_files.pop(path_key, None)
            continue

        manifest_files[path_key] = entry
        changed_chunk_count += len(entry.get("chunk_ids", []))
        if is_update:
            updated_file_count += 1
        else:
            new_file_count += 1

    manifest["embedding_model"] = EMBEDDING_MODEL
    manifest["collection_name"] = COLLECTION_NAME
    manifest["data_dirs"] = [str(path) for path in DATA_DIRS]
    save_manifest(manifest)

    elapsed = time.perf_counter() - start_time

    print("\n索引更新完成")
    print(f"新增檔案數：{new_file_count}")
    print(f"更新檔案數：{updated_file_count}")
    print(f"刪除檔案數：{deleted_file_count}")
    print(f"跳過未變更檔案數：{skipped_unchanged_count}")
    print(f"Chunks 新增/更新數：{changed_chunk_count}")
    print(f"耗時：{elapsed:.1f} 秒")


def keyword_score(question: str, document: str, metadata: dict) -> float:
    tokens = [
        token.lower()
        for token in re.split(r"[\s/\\_:.\-()[\]{}<>#，。；：、,;]+", question)
        if len(token.strip()) >= 2
    ]
    if not tokens:
        return 0.0

    document_lower = document.lower()
    file_name = str(metadata.get("file_name", "")).lower()
    function_or_class = str(metadata.get("function_or_class", "")).lower()
    section_title = str(metadata.get("section_title", "")).lower()

    score = 0.0
    max_score = len(tokens) * 4.5

    for token in tokens:
        if token in document_lower:
            score += 1.0
        if token in file_name:
            score += 1.5
        if token in function_or_class:
            score += 2.0
        if token in section_title:
            score += 2.0

    return min(score / max_score, 1.0)


def _vector_rank_score(rank: int, total: int) -> float:
    if total <= 1:
        return 1.0
    return max(0.0, 1.0 - (rank / (total - 1)))


def retrieve_docs(question: str) -> list[dict]:
    total_docs = collection.count()
    if total_docs == 0:
        return []

    query_embedding = encode_text(question)
    results = collection.query(
        query_embeddings=[query_embedding],
        n_results=min(VECTOR_SEARCH_TOP_K, total_docs),
        include=["documents", "metadatas", "distances"],
    )

    documents = results.get("documents", [[]])[0]
    metadatas = results.get("metadatas", [[]])[0]
    candidates: list[dict] = []

    for rank, (document, metadata) in enumerate(zip(documents, metadatas)):
        kw_score = keyword_score(question, document, metadata)
        candidates.append(
            {
                "document": document,
                "metadata": metadata,
                "rerank_score": None,
                "keyword_score": kw_score,
                "vector_rank_score": _vector_rank_score(rank, len(documents)),
                "final_score": 0.0,
            }
        )

    if not candidates:
        return []

    if reranker_model is not None:
        pairs = [[question, candidate["document"]] for candidate in candidates]
        try:
            scores = reranker_model.predict(pairs)
            for candidate, score in zip(candidates, scores):
                rerank_score = float(score)
                candidate["rerank_score"] = rerank_score
                candidate["final_score"] = rerank_score + 0.15 * candidate["keyword_score"]
        except Exception as exc:
            print(f"Reranker 評分失敗，改用向量搜尋結果：{exc}")
            for candidate in candidates:
                candidate["final_score"] = (
                    candidate["vector_rank_score"] + 0.25 * candidate["keyword_score"]
                )
    else:
        for candidate in candidates:
            candidate["final_score"] = (
                candidate["vector_rank_score"] + 0.25 * candidate["keyword_score"]
            )

    candidates.sort(key=lambda item: item["final_score"], reverse=True)
    return candidates[:RERANK_TOP_K]


def format_context(results: list[dict]) -> str:
    if not results:
        return "目前沒有可用的索引內容，請先建立或重建索引。"

    contexts: list[str] = []

    for result in results:
        doc = result["document"]
        meta = result["metadata"]
        rerank_score = result.get("rerank_score")
        labels = [
            f"來源：{meta.get('source_path', '')}",
            f"檔名：{meta.get('file_name', '')}",
            f"類型：{meta.get('file_type', '')}",
            f"資料夾：{meta.get('folder_name', '')}",
            f"chunk：{meta.get('chunk_id', '')}",
            f"Chunk策略：{meta.get('chunk_strategy', '')}",
            f"Rerank分數：{rerank_score:.4f}" if rerank_score is not None else "Rerank分數：N/A",
            f"Keyword分數：{result.get('keyword_score', 0.0):.4f}",
            f"Final分數：{result.get('final_score', 0.0):.4f}",
        ]

        if "page_number" in meta:
            labels.append(f"頁碼：{meta['page_number']}")
        if "section_title" in meta:
            labels.append(f"章節：{meta['section_title']}")
        if "function_or_class" in meta:
            labels.append(f"函式/類別：{meta['function_or_class']}")

        contexts.append(f"{' | '.join(labels)}\n內容：{doc}")

    return "\n\n---\n\n".join(contexts)


def search_docs(question: str) -> str:
    return format_context(retrieve_docs(question))


def ask_ollama(question: str, context: str, history: list[dict] | None = None) -> str:
    history_text = "無"
    if history:
        history_text = "\n\n".join(
            f"Q: {item.get('question', '')}\nA: {item.get('answer', '')}"
            for item in history
        )

    prompt = f"""
你是一個本地 RAG Agent。

規則：
1. 只能根據「檢索到的資料」回答。
2. 如果資料中沒有答案，請回答「資料中沒有提到」。
3. 請用繁體中文回答。
4. 回答時列出依據來源。
5. 對話紀錄只能作為理解上下文使用。
6. 最終回答仍必須以檢索內容為主要依據。
7. 如果檢索內容沒有答案，不能只靠記憶亂答。
8. 若問題是接續問題，例如「那它用什麼模型？」，可以參考最近對話補全主詞。

最近對話紀錄：
{history_text}

檢索到的資料：
{context}

問題：
{question}

回答：
"""

    response = requests.post(
        OLLAMA_URL,
        json={
            "model": MODEL,
            "prompt": prompt,
            "stream": False,
        },
        timeout=300,
    )

    response.raise_for_status()
    return response.json()["response"]


def main() -> None:
    while True:
        print("\n=== NKUST Local RAG Agent v0.6 ===")
        print("1. 增量更新索引")
        print("2. 完整重建索引")
        print("3. 進入常駐問答模式")
        print("q. 離開")

        mode = input("請選擇模式：").strip()

        if mode == "1":
            build_index(full_rebuild=False)
            print("\n索引增量更新完畢，已回到主選單。")
            continue

        elif mode == "2":
            build_index(full_rebuild=True)
            print("\n索引完整重建完畢，已回到主選單。")
            continue

        elif mode == "3":
            history = deque(maxlen=CONVERSATION_MEMORY_SIZE)

            print("\n已進入常駐問答模式。")
            print("輸入 back 可回到主選單。")
            print("輸入 rebuild 可增量更新索引。")
            print("輸入 full_rebuild 可完整重建索引。")
            print("輸入 memory 可查看最近對話。")
            print("輸入 clear_memory 可清空對話記憶。")
            print("輸入 q 可離開。\n")

            while True:
                question = input("\n請輸入問題：").strip()

                if question.lower() == "q":
                    print("已離開 Agent。")
                    return

                if question.lower() == "back":
                    print("已回到主選單。")
                    break

                if question.lower() == "rebuild":
                    build_index(full_rebuild=False)
                    print("\n索引增量更新完畢，回到問答模式。")
                    continue

                if question.lower() == "full_rebuild":
                    build_index(full_rebuild=True)
                    print("\n索引完整重建完畢，回到問答模式。")
                    continue

                if question.lower() == "memory":
                    if not history:
                        print("目前沒有對話記憶。")
                    else:
                        print("\n=== 最近對話記憶 ===\n")
                        for index, item in enumerate(history, start=1):
                            print(f"{index}. Q: {item.get('question', '')}")
                            print(f"   A: {item.get('answer', '')}\n")
                    continue

                if question.lower() == "clear_memory":
                    history.clear()
                    print("已清空對話記憶。")
                    continue

                if not question:
                    continue

                try:
                    context = search_docs(question)
                    answer = ask_ollama(question, context, list(history))
                except Exception as exc:
                    print(f"問答時發生錯誤：{exc}")
                    continue

                history.append(
                    {
                        "question": question,
                        "answer": answer,
                    }
                )

                print("\n=== 檢索內容 ===\n")
                print(context)

                print("\n=== Agent 回答 ===\n")
                print(answer)

        elif mode.lower() == "q":
            print("已離開 Agent。")
            break

        else:
            print("無效模式，請重新選擇。")


if __name__ == "__main__":
    main()
